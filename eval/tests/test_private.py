"""Private data (ADR 0014): own chat messages stay in data/private, and nothing derived from them is public."""

import json

import pytest

from arabizi_eval import b0, parity, run
from arabizi_eval import pairs as P
from arabizi_eval.pairs_cli import held_out_quota, main


@pytest.fixture
def folders(monkeypatch, tmp_path):
    public, private = tmp_path / "pairs", tmp_path / "private" / "pairs"
    public.mkdir()
    monkeypatch.setattr(P, "PAIRS_DIR", public)
    monkeypatch.setattr(P, "PRIVATE_DIR", private)
    monkeypatch.setattr(P, "MANIFEST_PATH", public / "manifest.json")
    monkeypatch.setattr(b0, "LEXICON_PATH", tmp_path / "lexicon.json")
    monkeypatch.setattr(parity, "FIXTURE", tmp_path / "b0_parity.json")  # `pairs lexicon` rewrites it
    monkeypatch.setattr(run, "RESULTS_DIR", tmp_path / "results")
    monkeypatch.setattr(run, "PRIVATE_RESULTS_DIR", tmp_path / "results" / "private")
    public_row = {"id": "ai-0001", "arabizi": "shu 3am ta3mel ya 7abibi", "arabic": "شو عم تعمل يا حبيبي",
                  "accepted": {}, "method": "ai_checked", "writer": "x", "checked_by": "moustafa", "source": "batch",
                  "split": "train", "date": "2026-10-13"}
    (public / "train-ai_checked.jsonl").write_text(json.dumps(public_row, ensure_ascii=False) + "\n", encoding="utf-8")
    chats = tmp_path / "chats.jsonl"
    lines = [{"arabizi": f"sirrewordzz {n} bas kif", "arabic": f"سرية {n} بس كيف", "source": "private chat 1",
              "year": 2025} for n in range(1, 31)]
    lines += [{"arabizi": f"wlek ya grp{n}", "arabic": f"ولك يا جماعة {n}", "source": "group chat"}
              for n in range(1, 11)]
    chats.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in lines) + "\n", encoding="utf-8")
    return tmp_path, public, private, chats


def test_held_out_quota_is_proportional_and_exact():
    assert held_out_quota({"a": 1017, "b": 891, "c": 79}, 300) == {"a": 154, "b": 134, "c": 12}
    assert sum(held_out_quota({"a": 3, "b": 3, "c": 3}, 4).values()) == 4


def test_import_chats_writes_only_to_the_private_folder(folders, capsys):
    tmp, public, private, chats = folders
    before = sorted(p.name for p in public.iterdir())
    assert main(["import-chats", str(chats), "--held-out", "10"]) == 0
    assert sorted(p.name for p in public.iterdir()) == before  # nothing new in the public folder
    rows = {s: [json.loads(x) for x in (private / f).read_text(encoding="utf-8").splitlines()]
            for s, f in (("test", "test.jsonl"), ("dev", "dev.jsonl"), ("train", "train-hand.jsonl"))}
    assert (len(rows["test"]), len(rows["dev"]), len(rows["train"])) == (7, 3, 30)
    groups = [r["source"] for r in rows["test"] + rows["dev"]]
    assert sum("group chat" in g for g in groups) == 2  # 10 of 40 rows -> 2 of 10 held out
    assert all(r["method"] == "hand" and r["writer"] == "moustafa" for r in rows["train"])
    assert "sirrewordzz" not in capsys.readouterr().out  # no message text printed
    assert P.validate().ok
    assert main(["import-chats", str(chats)]) == 1  # refuses to import twice


def test_variants_of_private_rows_stay_private(folders):
    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    assert main(["variants", "--per-row", "1", "--show", "0"]) == 0
    public_text = (public / "train-rule_variant.jsonl").read_text(encoding="utf-8")
    private_text = (private / "train-rule_variant.jsonl").read_text(encoding="utf-8")
    assert "sirrewordzz" not in public_text and "sirrewordzz" in private_text.lower()
    ids = [json.loads(x)["id"] for x in (public_text + private_text).splitlines()]
    assert len(ids) == len(set(ids))
    assert P.validate().ok


def test_public_dictionary_never_learns_private_words(folders):
    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    assert main(["lexicon"]) == 0
    lexicon = (tmp / "lexicon.json").read_text(encoding="utf-8")
    assert "sirrewordzz" not in lexicon and "shu" in json.loads(lexicon)["words"]
    assert "sirrewordzz" not in (tmp / "b0_parity.json").read_text(encoding="utf-8")  # the Swift fixture is public too


def test_runs_with_private_rows_go_to_the_private_results_folder(folders):
    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    assert run.main(["run", "--system", "identity", "--split", "dev"]) == 0
    assert run.main(["run", "--system", "identity", "--split", "dev", "--subset", "public"]) == 1  # no public dev
    (found,) = (tmp / "results" / "private").iterdir()
    meta = json.loads((found / "metrics.json").read_text(encoding="utf-8"))
    assert meta["private_sentences"] == 3 and not (tmp / "results" / found.name).exists()


def test_import_chats_with_nothing_importable_reports_instead_of_crashing(folders, capsys):
    tmp, public, private, chats = folders
    empty = tmp / "empty.jsonl"
    empty.write_text(json.dumps({"arabizi": "shu 3am ta3mel ya 7abibi", "arabic": "شو عم تعمل يا حبيبي"}) + "\n",
                     encoding="utf-8")  # repeats the public training row
    assert main(["import-chats", str(empty)]) == 1
    assert "every row was skipped" in capsys.readouterr().out


def test_publish_moves_chosen_rows_and_hides_the_chat_label(folders, monkeypatch, capsys):
    from arabizi_eval import pairs_cli

    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    monkeypatch.setattr(parity, "write", lambda: None)  # the fixture is tested on its own
    held_out_before = [r for r in pairs_cli.private_rows() if r["split"] in ("test", "dev")]
    answers = iter(["y", "n", "y", "q"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))
    assert main(["publish"]) == 0

    moved = [r for name in ("test.jsonl", "dev.jsonl") if (public / name).exists()
             for _, r, _ in P.read_jsonl(public / name)]
    assert len(moved) == 2
    assert all(r["source"] == pairs_cli.PUBLISHED_SOURCE for r in moved)  # no chat label in the public files
    still_private = {r["id"] for r in pairs_cli.private_rows()}
    assert not still_private & {r["id"] for r in moved}
    assert len(still_private & {r["id"] for r in held_out_before}) == len(held_out_before) - 2
    assert P.validate().ok  # ids unique, splits and methods still valid

    declined = json.loads(pairs_cli.publish_review_path().read_text(encoding="utf-8"))["declined"]
    assert len(declined) == 1 and pairs_cli.publish_review_path().is_relative_to(tmp)  # remembered, kept private
    capsys.readouterr()
    monkeypatch.setattr("builtins.input", lambda _prompt: "q")
    main(["publish"])
    # The next run offers what is left: neither the 2 published rows nor the declined one.
    assert f"{len(held_out_before) - 3} private test/dev rows to review" in capsys.readouterr().out


def test_publish_refuses_after_the_freeze(folders, capsys):
    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    P.MANIFEST_PATH.write_text("{}", encoding="utf-8")
    assert main(["publish"]) == 1
    assert "frozen" in capsys.readouterr().out


def test_publish_through_a_spreadsheet(folders, monkeypatch):
    import openpyxl

    from arabizi_eval import pairs_cli

    tmp, public, private, chats = folders
    main(["import-chats", str(chats), "--held-out", "10"])
    monkeypatch.setattr(parity, "write", lambda: None)
    assert main(["publish", "--sheet"]) == 0
    sheet_path = pairs_cli.publish_sheet_path()
    assert sheet_path.is_relative_to(tmp)  # written next to the private data, never in the public folder

    book = openpyxl.load_workbook(sheet_path)
    sheet = book.active
    assert [c.value for c in sheet[1]] == ["publish", "split", "arabizi", "arabic", "id"]
    marked = {sheet.cell(2, 5).value, sheet.cell(3, 5).value}
    sheet.cell(2, 1, "y")
    sheet.cell(3, 1, "Y")
    sheet.cell(4, 1, "n")
    sheet.cell(5, 1, "y")
    sheet.cell(5, 4, sheet.cell(5, 4).value + " edited")  # text changed in the sheet: not published
    book.save(sheet_path)

    assert main(["publish", "--apply", str(sheet_path)]) == 1  # 1: one row was refused
    moved = {r["id"] for name in ("test.jsonl", "dev.jsonl") if (public / name).exists()
             for _, r, _ in P.read_jsonl(public / name)}
    assert moved == marked
    declined = json.loads(pairs_cli.publish_review_path().read_text(encoding="utf-8"))["declined"]
    assert declined == [sheet.cell(4, 5).value]
    assert P.validate().ok
