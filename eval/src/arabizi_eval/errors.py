"""Error analysis of a scored run: every wrong word in a spreadsheet (roadmap P7 step 1, started early).

  uv run --project eval evalrun errors eval/results/<run_id>

Each wrong word gets a row: the reference word, the system's word, and a type suggested from the words alone
(English converted, MSA word, one letter off, joined/split...). Correct the type by hand where the suggestion is
wrong; "different word" rows need a person to say why (rare word, ambiguous word, reference wrong...). The summary
sheet counts the types as they are edited. A spreadsheet because terminals here cannot display Arabic.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from arabizi_eval import metrics

# Types a row can have; the first ones are suggested automatically, the last ones only a person can tell.
SUGGESTED = ("English converted", "left in Latin", "number", "MSA word", "spelling convention", "one letter off",
             "joined/split", "dropped word", "added word", "different word")
BY_HAND = ("rare word", "ambiguous word", "reference wrong", "other")
TYPES = SUGGESTED + BY_HAND


@dataclass(frozen=True)
class WordError:
    reference: str  # empty when the system added a word
    output: str  # empty when the system dropped a word
    type: str


def _matches(output: list[str], reference: list[frozenset[str]]) -> list[tuple[int, int]]:
    """(output index, reference index) of each matched word, from the same LCS that word accuracy counts."""
    table = [[0] * (len(reference) + 1) for _ in range(len(output) + 1)]
    for i, word in enumerate(output):
        for j, allowed in enumerate(reference):
            table[i + 1][j + 1] = table[i][j] + 1 if word in allowed else max(table[i][j + 1], table[i + 1][j])
    pairs, i, j = [], len(output), len(reference)
    while i and j:
        if output[i - 1] in reference[j - 1] and table[i][j] == table[i - 1][j - 1] + 1:
            pairs.append((i - 1, j - 1))
            i, j = i - 1, j - 1
        elif table[i - 1][j] >= table[i][j - 1]:
            i -= 1
        else:
            j -= 1
    return pairs[::-1]


def suggest_type(reference: str, output: str) -> str:
    if not output:
        return "dropped word"
    if not reference:
        return "added word"
    ref_latin, out_latin = bool(metrics.LATIN_WORD.search(reference)), bool(metrics.LATIN_WORD.search(output))
    if ref_latin and not out_latin:
        return "English converted"
    if out_latin and not ref_latin:
        return "left in Latin"
    if any(c.isdigit() for c in reference + output):
        return "number"
    if metrics.norm(output, True) == metrics.norm(reference, True):  # before MSA: اين for أين is spelling only
        return "spelling convention"
    if metrics.norm(output, True) in metrics.msa_words():
        return "MSA word"
    if metrics.edit_distance(metrics.norm(output, False), metrics.norm(reference, False)) <= 1:
        return "one letter off"
    return "different word"


def _gap_errors(ref_gap: list[str], out_gap: list[str]) -> list[WordError]:
    """Errors between two matched words: one joined/split error, or word by word in order."""
    if ref_gap and out_gap and len(ref_gap) != len(out_gap) and (
            "".join(metrics.norm(w, True) for w in ref_gap) == "".join(metrics.norm(w, True) for w in out_gap)):
        return [WordError(" ".join(ref_gap), " ".join(out_gap), "joined/split")]
    errors = []
    for k in range(max(len(ref_gap), len(out_gap))):
        ref = ref_gap[k] if k < len(ref_gap) else ""
        out = out_gap[k] if k < len(out_gap) else ""
        errors.append(WordError(ref, out, suggest_type(ref, out)))
    return errors


def sentence_errors(row: dict, prediction: str) -> list[WordError]:
    """Wrong words of one sentence, aligned against the reference that gives the best word accuracy."""
    accepted = row.get("accepted", {})
    out_words = metrics.tokens(prediction)
    out_norm = [metrics.norm(w, False) for w in out_words]
    texts = metrics.reference_texts(row["arabic"], accepted)
    refs = metrics.references(row["arabic"], accepted, lenient=False)
    best = max(range(len(refs)),
               key=lambda k: metrics.lcs(out_norm, refs[k]) / max(len(refs[k]), len(out_norm), 1))
    ref_words = metrics.tokens(texts[best])
    errors, i, j = [], 0, 0
    for out_index, ref_index in [*_matches(out_norm, refs[best]), (len(out_words), len(ref_words))]:
        errors += _gap_errors(ref_words[j:ref_index], out_words[i:out_index])
        i, j = out_index + 1, ref_index + 1
    return errors


def _append_text(sheet, values: list) -> None:
    """Append a row of data as plain text: no formulas (`=)` in a chat) and no control characters."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values])
    for cell in sheet[sheet.max_row]:
        if cell.data_type == "f":
            cell.data_type = "s"


def write_sheet(meta: dict, records: list[dict], rows: dict[str, dict], path: Path,
                warnings: list[str] | None = None) -> tuple[int, int]:
    """Write the error sheet; returns (error rows, sentences with errors).

    Rows are not the word-accuracy error count: a gap with one dropped and one added word is two rows but one
    wrong word in the score, and a joined/split error is one row for several words. Both are in the summary.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    book = openpyxl.Workbook()
    words = book.active
    words.title = "words"
    words.append(["id", "reference word", "output word", "type", "note", "arabizi", "reference", "output"])
    sentences = book.create_sheet("sentences")
    sentences.append(["id", "words right", "words total", "arabizi", "reference", "output"])
    n_rows = n_sentences = 0
    for record in records:
        row = rows.get(record["id"], {"arabic": record["reference"]})
        found = sentence_errors(row, record["prediction"])
        if not found:
            continue
        n_sentences += 1
        _append_text(sentences, [record["id"], record["words_correct"], record["words_total"], record["arabizi"],
                                 record["reference"], record["prediction"]])
        for error in found:
            n_rows += 1
            _append_text(words, [record["id"], error.reference, error.output, error.type, "", record["arabizi"],
                                 record["reference"], record["prediction"]])

    choice = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=True,
                            errorStyle="warning", error="Not one of the listed types; keep it anyway?")
    words.add_data_validation(choice)
    choice.add(f"D2:D{max(n_rows + 1, 2)}")
    right_to_left = Alignment(horizontal="right", readingOrder=2)
    for sheet, arabic, widths in ((words, "BCGH", (10, 18, 18, 20, 30, 50, 50, 50)),
                                  (sentences, "EF", (10, 7, 7, 50, 50, 50))):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in arabic:
            for cell in sheet[column][1:]:
                cell.alignment = right_to_left
        for index, width in enumerate(widths):
            sheet.column_dimensions[chr(ord("A") + index)].width = width
        sheet.freeze_panes = "A2"

    summary = book.create_sheet("summary", 0)
    m = meta.get("metrics", {})
    wrong = sum(r["words_total"] - r["words_correct"] for r in records)
    labels = (("run", meta.get("run_id")), ("split", f"{meta.get('split')} ({meta.get('subset')})"),
              ("sentences", m.get("sentences")), ("word accuracy", m.get("word_accuracy")),
              ("sentences with errors", n_sentences), ("wrong words (word accuracy)", wrong),
              ("error rows in this sheet", n_rows))
    for label, value in labels:
        summary.append([label, value])
    for warning in warnings or []:
        summary.append(["warning", warning])
    summary.append([])
    summary.append(["type", "rows", "share"])
    header = summary.max_row
    for row_number, name in enumerate(TYPES, header + 1):
        summary.append([name, f'=COUNTIF(words!D:D,A{row_number})',
                        f"=IF($B${len(TYPES) + header + 1}=0,0,B{row_number}/$B${len(TYPES) + header + 1})"])
    total_row = summary.max_row + 1
    summary.append(["total", f"=SUM(B{header + 1}:B{total_row - 1})", ""])
    for cell in (*summary["A"][:header - 2], *summary[header], *summary[total_row]):
        cell.font = Font(bold=True)
    for cell in summary["C"][header:]:
        cell.number_format = "0%"
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 50
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return n_rows, n_sentences
